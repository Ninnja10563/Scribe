#!/usr/bin/env python3
"""Independently inspect macOS smoke artifacts. Development-only dependencies:
python-docx==1.2.0 and pymupdf==1.28.2. This is not a Word compatibility certification.
"""
import argparse
import json
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree

from docx import Document
import pymupdf

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('directory', type=Path, help='Downloaded artifact smoke/ directory')
args = parser.parse_args()
root = args.directory
native = json.loads((root / 'Smoke.scribe').read_text())
word = Document(root / 'Smoke.docx')
assert len(word.tables) == 1 and len(word.inline_shapes) == 1, 'Missing structured objects'
assert word.tables[0].cell(0, 0).text == 'Section'
assert word.tables[0].cell(2, 2).text == 'Ready'

namespace = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
value_key = '{' + namespace['w'] + '}val'
with ZipFile(root / 'Smoke.docx') as package:
    numbering = ElementTree.fromstring(package.read('word/numbering.xml'))
    starts = {node.get(value_key) for node in numbering.findall('.//w:start', namespace)}
    formats = {node.get(value_key) for node in numbering.findall('.//w:numFmt', namespace)}
    assert {'4', '9'} <= starts
    assert {'upperRoman', 'lowerLetter'} <= formats
    body = ElementTree.fromstring(package.read('word/document.xml'))
    if native.get('formatVersion', 0) >= 5:
        bookmark_names = {node.get('{' + namespace['w'] + '}name') for node in body.findall('.//w:bookmarkStart', namespace)}
        internal_links = [node.get('{' + namespace['w'] + '}anchor') for node in body.findall('.//w:hyperlink', namespace)]
        assert any(anchor in bookmark_names for anchor in internal_links), 'Missing Word internal destination'
    if native.get('formatVersion', 0) >= 6:
        assert len(native.get('tablesOfContents', [])) == 1
        instructions = [node.text or '' for node in body.findall('.//w:instrText', namespace)]
        assert any(code.strip().startswith('TOC ') for code in instructions), 'Missing actual Word TOC field'
        entries = [p for section in native['sections'] for p in section['paragraphs'] if p.get('toc', {}).get('kind') == 'entry']
        assert entries and all('—' not in ''.join(run['text'] for run in entry['runs']) for entry in entries), 'Unresolved TOC page labels'
    comments = native.get('comments', [])
    assert len(word.comments) == len(comments), 'Missing review text'
    for actual, expected in zip(word.comments, comments):
        assert actual.text == expected['text'] and actual.author == expected['author']
    if comments:
        comment_xml = ElementTree.fromstring(package.read('word/comments.xml'))
        extended = ElementTree.fromstring(package.read('word/commentsExtended.xml'))
        w14 = '{http://schemas.microsoft.com/office/word/2010/wordml}'
        w15 = '{http://schemas.microsoft.com/office/word/2012/wordml}'
        states = {item.get(w15 + 'paraId'): item.get(w15 + 'done') in ('1', 'true', 'on') for item in extended}
        for element, expected in zip(comment_xml, comments):
            final_paragraph = element.findall('w:p', namespace)[-1]
            assert states[final_paragraph.get(w14 + 'paraId')] == expected['resolved']
    attached_count = sum(not comment.get('isDetached', False) for comment in comments)
    for tag in ('commentRangeStart', 'commentRangeEnd', 'commentReference'):
        assert len(body.findall('.//w:' + tag, namespace)) == attached_count

with pymupdf.open(root / 'Smoke.pdf') as full, pymupdf.open(root / 'Selected-pages.pdf') as selected:
    assert len(full) >= 10 and len(selected) == 2
    assert selected[0].get_text() == full[0].get_text()
    assert selected[1].get_text() == full[-1].get_text()
    assert selected.metadata['title'] == 'Selected pages'
    assert selected.metadata['subject'] == 'Range export'
    if native.get('formatVersion', 0) >= 5:
        for pdf in (full, selected):
            links = pdf[-1].get_links()
            assert any(link.get('kind') == pymupdf.LINK_GOTO and link.get('page') == 0 for link in links), 'Missing PDF internal destination'
            assert not any(link.get('uri', '').startswith('scribe:') for link in links), 'Private application URL leaked into PDF'
    if native.get('formatVersion', 0) >= 6:
        spans = [span for block in full[0].get_text('dict')['blocks'] for line in block.get('lines', []) for span in line['spans']]
        toc_links = [span for span in spans if 'A considered place to write' in span['text'] and span['size'] < 20]
        assert toc_links and all((span['color'] & 255) > ((span['color'] >> 16) & 255) for span in toc_links), 'PDF links lost their editor styling'
    text = '\n'.join(page.get_text() for page in full)
    for marker in ('IV.', 'a.', 'V.', 'IX.', 'Paragraph 80.', f'Page {len(full)} of {len(full)}'):
        assert marker in text, f'Missing output: {marker}'
    for comment in comments:
        assert comment['text'] not in text, 'Review sidebar text leaked into printed body'
    print(f'Verified {len(full)} source pages, two selected pages, metadata, numbering, table, image and {len(comments)} comments.')
