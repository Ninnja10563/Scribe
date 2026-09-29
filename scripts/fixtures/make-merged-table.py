"""Original merged-cell fixture; python-docx 1.2.0, not Word-produced."""
from pathlib import Path
from docx import Document

doc = Document()
doc.add_paragraph('Independent merged table fixture')
table = doc.add_table(rows=4, cols=3)
for row in range(4):
    for column in range(3):
        table.cell(row, column).text = f'Cell {row},{column}: café 東京'
table.cell(0, 0).merge(table.cell(1, 1))
table.cell(2, 1).merge(table.cell(3, 1))
doc.save(Path(__file__).resolve().parents[2] / 'Tests/ImportExportTests/Fixtures/MergedTable.docx')
