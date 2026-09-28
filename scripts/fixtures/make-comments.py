"""Original multi-paragraph, overlapping-comment fixture; requires python-docx 1.2.0."""
from pathlib import Path
from docx import Document

doc = Document()
first = doc.add_paragraph().add_run('First paragraph')
second = doc.add_paragraph().add_run('Second paragraph')
doc.add_comment(runs=[first, second], text='Across both paragraphs\nKeep Unicode: café 東京', author='Independent editor', initials='IE')
doc.add_comment(runs=first, text='Overlapping first paragraph', author='Second editor', initials='SE')
doc.save(Path(__file__).resolve().parents[2] / 'Tests/ImportExportTests/Fixtures/Comments.docx')
