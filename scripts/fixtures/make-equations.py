#!/usr/bin/env python3
"""Original test fixture authored with python-docx/lxml; not a Word-produced file."""
from pathlib import Path
from docx import Document
from lxml import etree

root = Path(__file__).resolve().parents[2]
document = Document()
paragraph = document.add_paragraph('Independent equation: ')
math = etree.fromstring(b'''<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
<m:f><m:num><m:r><m:t>12</m:t></m:r></m:num><m:den><m:sSubSup><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sub><m:r><m:t>i</m:t></m:r></m:sub><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSubSup></m:den></m:f>
</m:oMath>''')
paragraph._p.append(math)
paragraph.add_run(' remains editable.')
document.save(root / 'Tests/ImportExportTests/Fixtures/IndependentEquations.docx')
