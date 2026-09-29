"""Original independent python-docx input with DrawingML image transforms."""
from pathlib import Path
from io import BytesIO
import base64
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement

source = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAACAAAAAQCAIAAAD4YuoOAAAAKElEQVR4nGP4z8BAEiJNNRlo1IJRCwbAAlJ1/P/PQBIatWDUgiFgAQAk0H2fCntH2wAAAABJRU5ErkJggg==')
document = Document()
document.core_properties.title = 'Independent adjusted image'
document.add_paragraph('Independent crop, clockwise rotation and opacity')
shape = document.add_picture(BytesIO(source), width=Pt(150), height=Pt(100))
picture = shape._inline.graphic.graphicData.pic
picture.spPr.xfrm.set('rot', '5400000')
crop = OxmlElement('a:srcRect'); crop.set('l', '25000')
picture.blipFill.insert(1, crop)
alpha = OxmlElement('a:alphaModFix'); alpha.set('amt', '50000')
picture.blipFill.blip.append(alpha)
document.save(Path(__file__).resolve().parents[2] / 'Tests/ImportExportTests/Fixtures/ImageAdjustments.docx')
