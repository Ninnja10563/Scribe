"""Original styled-cell fixture; requires python-docx 1.2.0. Not Word-produced."""
from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

doc = Document()
doc.add_paragraph('Independent cell formatting fixture')
table = doc.add_table(rows=2, cols=2)
for row in range(2):
    for column in range(2):
        table.cell(row, column).text = f'Cell {row},{column}: café 東京'
first = table.cell(0, 0)
shade = OxmlElement('w:shd'); shade.set(qn('w:val'), 'clear'); shade.set(qn('w:fill'), 'DDEEFF')
first._tc.get_or_add_tcPr().append(shade)
first.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
last = table.cell(1, 1)
properties = last._tc.get_or_add_tcPr()
borders = OxmlElement('w:tcBorders')
for edge in ('top', 'left', 'bottom', 'right'):
    element = OxmlElement('w:' + edge)
    for key, value in (('val', 'single'), ('sz', '12'), ('color', '336699')):
        element.set(qn('w:' + key), value)
    borders.append(element)
properties.append(borders)
margins = OxmlElement('w:tcMar')
for edge in ('top', 'left', 'bottom', 'right'):
    element = OxmlElement('w:' + edge); element.set(qn('w:w'), '160'); element.set(qn('w:type'), 'dxa'); margins.append(element)
properties.append(margins)
last.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
table.rows[1].height = Pt(90)
table.rows[1].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
path = Path(__file__).resolve().parents[2] / 'Tests/ImportExportTests/Fixtures/TableFormatting.docx'
doc.save(path)
