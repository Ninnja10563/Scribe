#!/usr/bin/env python3
"""Original first/even running-content fixture; requires python-docx 1.2.0."""
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

document = Document()
document.settings.odd_and_even_pages_header_footer = True
section = document.sections[0]
section.different_first_page_header_footer = True
section.header.paragraphs[0].text = 'Independent default header'
section.footer.paragraphs[0].text = 'Independent default footer'
section.first_page_header.paragraphs[0].text = ''
section.first_page_footer.paragraphs[0].text = 'Cover footer — résumé'
section.even_page_header.paragraphs[0].text = 'Independent even header'
section.even_page_footer.paragraphs[0].text = 'Independent even footer'
numbering = OxmlElement('w:pgNumType'); numbering.set(qn('w:start'), '2')
section._sectPr.insert_element_before(numbering, 'w:cols', 'w:titlePg')
for index in range(3):
    if index:
        document.add_page_break()
    document.add_paragraph(f'Independent page {index + 1}')
destination = Path(__file__).resolve().parents[2] / 'Tests/ImportExportTests/Fixtures/RunningContent.docx'
document.save(destination)
print(destination)
