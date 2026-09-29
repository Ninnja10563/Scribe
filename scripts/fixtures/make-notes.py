#!/usr/bin/env python3
"""Independent OOXML fixture built with python-docx/lxml, not Microsoft Word."""
from pathlib import Path
from docx import Document
from docx.opc.part import Part
from lxml import etree

root = Path(__file__).resolve().parents[2]
word = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
relations = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
document = Document()
for kind, number, text in [('footnote', 17, 'Independent footnote résumé.'), ('endnote', 42, 'Independent endnote source.')]:
    paragraph = document.add_paragraph(f'Claim with {kind} ')
    reference = etree.Element(f'{{{word}}}{kind}Reference'); reference.set(f'{{{word}}}id', str(number))
    paragraph.add_run()._r.append(reference)
    part_name = f'/word/custom-{kind}s.xml'
    xml = f'''<q:{kind}s xmlns:q="{word}">
<q:{kind} q:type="separator" q:id="-1"><q:p><q:r><q:separator/></q:r></q:p></q:{kind}>
<q:{kind} q:id="{number}"><q:p><q:r><q:{kind}Ref/></q:r><q:r><q:t xml:space="preserve"> </q:t></q:r><q:r><q:rPr><q:i/></q:rPr><q:t>{text}</q:t></q:r></q:p></q:{kind}>
</q:{kind}s>'''
    from docx.opc.packuri import PackURI
    part = Part(PackURI(part_name), f'application/vnd.openxmlformats-officedocument.wordprocessingml.{kind}s+xml', xml.encode(), document.part.package)
    document.part.relate_to(part, relations + kind + 's')
document.save(root / 'Tests/ImportExportTests/Fixtures/IndependentNotes.docx')
