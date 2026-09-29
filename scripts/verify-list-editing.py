#!/usr/bin/env python3
"""Validate exports from the native ordinary list-boundary deletion test."""
import json
import re
import sys
from pathlib import Path
import docx
import pymupdf

folder = Path(sys.argv[1])
source = json.loads((folder / "JoinedList.scribe").read_text())
paragraphs = source["sections"][0]["paragraphs"]
assert len(paragraphs) == 1
assert "".join(run["text"] for run in paragraphs[0]["runs"]) == "FirstSecond"
assert paragraphs[0]["list"]["start"] == 4
assert all(not run.get("review") for run in paragraphs[0]["runs"])
word = docx.Document(folder / "JoinedList.docx")
assert [p.text for p in word.paragraphs] == ["FirstSecond"]
assert word.paragraphs[0]._p.pPr.numPr is not None
with pymupdf.open(folder / "JoinedList.pdf") as pdf:
    assert len(pdf) == 1
    text = pdf[0].get_text()
    assert text.count("FirstSecond") == 1
    assert len(re.findall(r"(?:^|\s)4\.", text)) == 1
    assert "5." not in text
    assert not pdf[0].get_images()
print("Native list join exports preserve authored text, one generated number and vector PDF text.")
